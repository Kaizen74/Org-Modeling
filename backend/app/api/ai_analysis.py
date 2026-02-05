"""AI Analysis API endpoints."""

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from typing import Optional, List, Union
from pydantic import BaseModel
import json
import os

from ..models.database import get_session
from ..models.models import OrgAnalysis, GradeSalary
from ..services.ai_analysis_service import AIAnalysisService
from ..services.metrics_service import MetricsCalculator
from ..services.claude_service import claude_service
from ..services.work_activities_service import WorkActivitiesAnalysisService

router = APIRouter(prefix="/api/v1/ai-analysis", tags=["ai-analysis"])


class AnalysisRequest(BaseModel):
    analysis_id: Optional[str] = None
    design_criteria: Optional[str] = None
    strategy_text: Optional[str] = None
    analysis_scope: str = "organization"  # "organization" or "department"
    department: Optional[str] = None  # Required if analysis_scope is "department"


@router.post("/analyze")
async def analyze_org(
    request: AnalysisRequest,
    db: AsyncSession = Depends(get_session)
):
    """
    Run AI analysis with structured categorization.

    Uses the latest org data if analysis_id is not provided.
    """
    # Get org analysis
    if request.analysis_id:
        result = await db.execute(
            select(OrgAnalysis).where(OrgAnalysis.id == request.analysis_id)
        )
        analysis = result.scalar_one_or_none()
    else:
        result = await db.execute(
            select(OrgAnalysis).order_by(OrgAnalysis.created_at.desc()).limit(1)
        )
        analysis = result.scalar_one_or_none()

    if not analysis:
        raise HTTPException(
            status_code=404,
            detail="No org data found. Please upload a CSV first."
        )

    employees = analysis.raw_data.get("employees", []) if analysis.raw_data else []

    if not employees:
        raise HTTPException(
            status_code=400,
            detail="No employee data found"
        )

    # Get metrics (calculate if not cached)
    if not analysis.metrics:
        # Load grade order
        grade_result = await db.execute(
            select(GradeSalary).order_by(GradeSalary.display_order)
        )
        grades = grade_result.scalars().all()
        grade_order = {g.grade: g.display_order for g in grades} if grades else {}

        calculator = MetricsCalculator(employees=employees, grade_order=grade_order)
        metrics = calculator.calculate_all_metrics()
        analysis.metrics = metrics
    else:
        metrics = analysis.metrics

    # Get API key from claude_service (handles .env file loading)
    if not claude_service.is_configured():
        raise HTTPException(
            status_code=400,
            detail="Claude API key not configured. Please set it in Settings."
        )

    # Run AI analysis
    service = AIAnalysisService(api_key=claude_service.api_key)

    strategy_docs = None
    if request.strategy_text:
        strategy_docs = [request.strategy_text]

    ai_result = await service.analyze_organization(
        metrics=metrics,
        employees=employees,
        strategy_documents=strategy_docs,
        design_criteria=request.design_criteria,
        analysis_scope=request.analysis_scope,
        department=request.department
    )

    # Store results
    analysis.ai_analysis = ai_result
    analysis.design_criteria = request.design_criteria
    analysis.strategy_context = request.strategy_text
    await db.commit()

    return {
        "analysis_id": analysis.id,
        "analysis_name": analysis.name,
        "ai_analysis": ai_result
    }


@router.post("/analyze-with-documents")
async def analyze_with_documents(
    analysis_id: Optional[str] = Form(default=None),
    design_criteria: Optional[str] = Form(default=None),
    analysis_scope: str = Form(default="organization"),
    department: Optional[str] = Form(default=None),
    strategy_docs: Union[List[UploadFile], UploadFile, None] = File(default=None),
    db: AsyncSession = Depends(get_session)
):
    """
    Run AI analysis with uploaded strategy documents (PDF/DOCX/PPTX).

    Accepts single or multiple file uploads.
    Supports organization-wide or department-level analysis.
    """
    # Normalize strategy_docs to always be a list
    if strategy_docs is None:
        docs_list: List[UploadFile] = []
    elif isinstance(strategy_docs, list):
        docs_list = strategy_docs
    else:
        docs_list = [strategy_docs]

    # Get org analysis
    if analysis_id:
        result = await db.execute(
            select(OrgAnalysis).where(OrgAnalysis.id == analysis_id)
        )
        analysis = result.scalar_one_or_none()
    else:
        result = await db.execute(
            select(OrgAnalysis).order_by(OrgAnalysis.created_at.desc()).limit(1)
        )
        analysis = result.scalar_one_or_none()

    if not analysis:
        raise HTTPException(
            status_code=404,
            detail="No org data found. Please upload a CSV first."
        )

    employees = analysis.raw_data.get("employees", []) if analysis.raw_data else []
    metrics = analysis.metrics

    if not metrics:
        # Try to calculate metrics if not available
        from ..models.models import GradeSalary
        grade_result = await db.execute(
            select(GradeSalary).order_by(GradeSalary.display_order)
        )
        grades = grade_result.scalars().all()
        grade_order = {g.grade: g.display_order for g in grades} if grades else {}

        # Enrich employees with salary from grade config
        grade_salary_map = {g.grade: g.median_salary for g in grades}
        enriched_employees = []
        for emp in employees:
            emp_copy = emp.copy()
            if not emp_copy.get("salary", 0):
                grade = emp_copy.get("grade", "")
                if grade in grade_salary_map:
                    emp_copy["salary"] = grade_salary_map[grade]
            enriched_employees.append(emp_copy)

        from ..services.metrics_service import MetricsCalculator
        calculator = MetricsCalculator(employees=enriched_employees, grade_order=grade_order)
        metrics = calculator.calculate_all_metrics()
        analysis.metrics = metrics
        await db.commit()

    # Extract document text
    doc_contents = []
    if docs_list:
        for doc in docs_list:
            if doc.filename:  # Skip empty uploads
                content = await _extract_text(doc)
                if content and not content.startswith("["):  # Skip error messages
                    doc_contents.append(f"--- {doc.filename} ---\n{content}")

    # Get API key from claude_service (handles .env file loading)
    if not claude_service.is_configured():
        raise HTTPException(
            status_code=400,
            detail="Claude API key not configured. Please set it in Settings."
        )

    # Run analysis
    service = AIAnalysisService(api_key=claude_service.api_key)

    ai_result = await service.analyze_organization(
        metrics=metrics,
        employees=employees,
        strategy_documents=doc_contents if doc_contents else None,
        design_criteria=design_criteria,
        analysis_scope=analysis_scope,
        department=department
    )

    # Store results
    analysis.ai_analysis = ai_result
    analysis.design_criteria = design_criteria
    if doc_contents:
        analysis.strategy_context = "\n\n".join(doc_contents)
    await db.commit()

    return {
        "analysis_id": analysis.id,
        "ai_analysis": ai_result
    }


@router.get("/quick-analysis")
async def get_quick_analysis(db: AsyncSession = Depends(get_session)):
    """
    Get a quick, non-AI analysis based on metrics alone.

    This doesn't require an API key.
    """
    result = await db.execute(
        select(OrgAnalysis).order_by(OrgAnalysis.created_at.desc()).limit(1)
    )
    analysis = result.scalar_one_or_none()

    if not analysis or not analysis.metrics:
        raise HTTPException(
            status_code=404,
            detail="No metrics available. Please upload and analyze org data first."
        )

    service = AIAnalysisService()
    quick = service.get_quick_analysis(analysis.metrics)

    return {
        "analysis_id": analysis.id,
        "analysis_name": analysis.name,
        **quick
    }


@router.get("/latest")
async def get_latest_ai_analysis(db: AsyncSession = Depends(get_session)):
    """Get the most recent AI analysis results."""
    result = await db.execute(
        select(OrgAnalysis)
        .where(OrgAnalysis.ai_analysis.isnot(None))
        .order_by(OrgAnalysis.created_at.desc())
        .limit(1)
    )
    analysis = result.scalar_one_or_none()

    if not analysis:
        raise HTTPException(
            status_code=404,
            detail="No AI analysis found. Please run an analysis first."
        )

    return {
        "analysis_id": analysis.id,
        "analysis_name": analysis.name,
        "ai_analysis": analysis.ai_analysis,
        "design_criteria": analysis.design_criteria,
        "created_at": analysis.created_at.isoformat() if analysis.created_at else None
    }


class WorkActivitiesRequest(BaseModel):
    analysis_id: Optional[str] = None
    industry: Optional[str] = None


@router.post("/work-activities")
async def analyze_work_activities(
    request: WorkActivitiesRequest,
    db: AsyncSession = Depends(get_session)
):
    """
    Run AI analysis on work activities to assess coherence, synergy, and themes.

    Analyzes:
    - Departmental coherence of work activities
    - Synergy between roles and teams
    - Work theme synthesis across the organization
    - Industry benchmark comparison
    """
    # Get org analysis
    if request.analysis_id:
        result = await db.execute(
            select(OrgAnalysis).where(OrgAnalysis.id == request.analysis_id)
        )
        analysis = result.scalar_one_or_none()
    else:
        result = await db.execute(
            select(OrgAnalysis).order_by(OrgAnalysis.created_at.desc()).limit(1)
        )
        analysis = result.scalar_one_or_none()

    if not analysis:
        raise HTTPException(
            status_code=404,
            detail="No org data found. Please upload a CSV first."
        )

    employees = analysis.raw_data.get("employees", []) if analysis.raw_data else []

    if not employees:
        raise HTTPException(
            status_code=400,
            detail="No employee data found"
        )

    # Check if any employees have work activities
    employees_with_activities = [
        emp for emp in employees
        if emp.get("work_activities", "").strip()
    ]

    if not employees_with_activities:
        raise HTTPException(
            status_code=400,
            detail="No work activities found in employee data. Please ensure your CSV includes a 'Work Activities' column."
        )

    # Get API key
    if not claude_service.is_configured():
        raise HTTPException(
            status_code=400,
            detail="Claude API key not configured. Please set it in Settings."
        )

    # Run work activities analysis
    service = WorkActivitiesAnalysisService(api_key=claude_service.api_key)

    work_activities_result = await service.analyze_work_activities(
        employees=employees,
        industry=request.industry
    )

    # Store results in the analysis
    if not analysis.ai_analysis:
        analysis.ai_analysis = {}

    analysis.ai_analysis["work_activities_analysis"] = work_activities_result
    await db.commit()

    return {
        "analysis_id": analysis.id,
        "analysis_name": analysis.name,
        "work_activities_analysis": work_activities_result
    }


@router.get("/work-activities/quick")
async def get_quick_work_activities_analysis(db: AsyncSession = Depends(get_session)):
    """
    Get quick work activities statistics without AI analysis.

    Returns coverage and basic breakdown by department.
    """
    result = await db.execute(
        select(OrgAnalysis).order_by(OrgAnalysis.created_at.desc()).limit(1)
    )
    analysis = result.scalar_one_or_none()

    if not analysis:
        raise HTTPException(
            status_code=404,
            detail="No org data found. Please upload a CSV first."
        )

    employees = analysis.raw_data.get("employees", []) if analysis.raw_data else []

    service = WorkActivitiesAnalysisService()
    quick = service.get_quick_analysis(employees)

    return {
        "analysis_id": analysis.id,
        "analysis_name": analysis.name,
        **quick
    }


@router.get("/work-activities/latest")
async def get_latest_work_activities_analysis(db: AsyncSession = Depends(get_session)):
    """Get the most recent work activities analysis results."""
    result = await db.execute(
        select(OrgAnalysis)
        .where(OrgAnalysis.ai_analysis.isnot(None))
        .order_by(OrgAnalysis.created_at.desc())
        .limit(1)
    )
    analysis = result.scalar_one_or_none()

    if not analysis:
        raise HTTPException(
            status_code=404,
            detail="No AI analysis found. Please run an analysis first."
        )

    work_activities = analysis.ai_analysis.get("work_activities_analysis") if analysis.ai_analysis else None

    if not work_activities:
        raise HTTPException(
            status_code=404,
            detail="No work activities analysis found. Please run a work activities analysis first."
        )

    return {
        "analysis_id": analysis.id,
        "analysis_name": analysis.name,
        "work_activities_analysis": work_activities,
        "created_at": analysis.created_at.isoformat() if analysis.created_at else None
    }


async def _extract_text(file: UploadFile) -> str:
    """Extract text from PDF, DOCX, PPTX, or TXT files."""
    content = await file.read()
    filename = file.filename.lower() if file.filename else ""

    try:
        if filename.endswith(".pdf"):
            import io
            try:
                import PyPDF2
                reader = PyPDF2.PdfReader(io.BytesIO(content))
                text = "\n".join([page.extract_text() for page in reader.pages])
                return text
            except ImportError:
                return "[PDF extraction requires PyPDF2]"

        elif filename.endswith(".docx"):
            import io
            try:
                import docx
                doc = docx.Document(io.BytesIO(content))
                text = "\n".join([para.text for para in doc.paragraphs])
                return text
            except ImportError:
                return "[DOCX extraction requires python-docx]"

        elif filename.endswith(".pptx"):
            import io
            try:
                from pptx import Presentation
                prs = Presentation(io.BytesIO(content))
                text_parts = []
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            text_parts.append(shape.text)
                return "\n".join(text_parts)
            except ImportError:
                return "[PowerPoint extraction requires python-pptx]"

        elif filename.endswith(".txt"):
            return content.decode("utf-8")

        else:
            return f"[Unsupported file type: {filename}]"

    except Exception as e:
        return f"[Error extracting text: {str(e)}]"
